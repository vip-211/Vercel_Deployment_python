# this is used for schheduling the email when it runs on linux server it will automaticaly send email to manager on daily basis sreport 
#comand to open linux scheduler  c- rontab -e
#Add this entry, replacing the example paths with your actual project and virtual-environment paths:     CRON_TZ=Asia/Kolkata
#                                                                                                         30 23 * * * cd /path/to/Deployment_Tracker && /path/to/venv/bin/python daily_deployment_report.py >> /path/to/Deployment_Tracker/logs/daily_report_cron.log 2>&1
import logging
import os
import re
import smtplib
import ssl
from datetime import date, datetime, time
from email.message import EmailMessage
from email.utils import formataddr
from pathlib import Path
from zoneinfo import ZoneInfo

import psycopg2
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

REPORT_TIMEZONE_NAME = os.getenv("REPORT_TIMEZONE", "Asia/Kolkata")
REPORT_TZ = ZoneInfo(REPORT_TIMEZONE_NAME)
REPORT_DIR = BASE_DIR / "reports"
REPORT_DIR.mkdir(parents=True, exist_ok=True)
LOG_DIR = BASE_DIR / "logs"
LOG_DIR.mkdir(parents=True, exist_ok=True)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    handlers=[
        logging.FileHandler(LOG_DIR / "daily_deployment_report.log", encoding="utf-8"),
        logging.StreamHandler(),
    ],
)
logger = logging.getLogger("daily_deployment_report")

EXPECTED_STATUSES = [
    "Pending", "Approved", "Rejected", "Scheduled", "In Progress",
    "Successful", "Failed", "Rolled Back", "Cancelled",
]

DETAIL_HEADERS = [
    "SR Number", "CR Number", "RFID", "Deployment Date", "Slot Start",
    "Slot End", "Application Name", "Deployment Type", "Change Type",
    "Status", "Change Requestor", "Released By", "Verified and Deployed By",
    "Incident Number", "Roll Back", "Last Modified By Email", "Last Modified At",
    "Release Notes", "Comments",
]


def get_connection():
    """Connect using the project's .env PostgreSQL settings."""
    required = ["DB_HOST", "DB_NAME", "DB_USER", "DB_PASSWORD"]
    missing = [key for key in required if not (os.getenv(key) or "").strip()]
    if missing:
        raise RuntimeError("Missing database setting(s) in .env: " + ", ".join(missing))

    return psycopg2.connect(
        host=os.getenv("DB_HOST"),
        port=int(os.getenv("DB_PORT", "5432")),
        dbname=os.getenv("DB_NAME"),
        user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASSWORD"),
        sslmode=os.getenv("DB_SSLMODE", "disable"),
        connect_timeout=15,
        application_name="Deployment Tracker Daily Report",
    )


def get_manager_recipients():
    """Only addresses explicitly listed in MANAGER_EMAILS can receive this report."""
    raw = os.getenv("MANAGER_EMAILS", "")
    recipients = sorted({
        address.strip().lower()
        for address in re.split(r"[,;\n]+", raw)
        if address.strip()
    })
    if not recipients:
        raise RuntimeError(
            "MANAGER_EMAILS is empty. Set only the manager email address(es) in .env."
        )

    invalid = [
        address for address in recipients
        if len(address) > 254 or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", address)
    ]
    if invalid:
        raise RuntimeError("Invalid email address(es) in MANAGER_EMAILS: " + ", ".join(invalid))
    return recipients


def fetch_daily_deployments(report_date):
    """Fetch the day's rows and the most recent recorded audit email per row."""
    sql = """
        SELECT
            d.id,
            d.cr_number,
            d.rfid,
            d.deployment_date,
            d.slot_start,
            d.slot_end,
            d.application_name,
            d.deployment_type,
            d.change_type,
            d.status,
            d.change_requestor,
            d.released_by,
            d.verified_and_deployed_by,
            d.incident_number,
            d.roll_back,
            latest.modified_by AS last_modified_by_email,
            latest.modified_at AS last_modified_at,
            d.release_notes,
            d.comments
        FROM public.deployment_tracker AS d
        LEFT JOIN LATERAL (
            SELECT h.modified_by, h.modified_at
            FROM public.deployment_change_history AS h
            WHERE h.deployment_id = d.id
            ORDER BY h.modified_at DESC, h.id DESC
            LIMIT 1
        ) AS latest ON TRUE
        WHERE d.deployment_date = %s
        ORDER BY d.slot_start NULLS LAST, d.id ASC
    """

    conn = None
    try:
        conn = get_connection()
        with conn.cursor() as cur:
            cur.execute(sql, (report_date,))
            columns = [description[0] for description in cur.description]
            return [dict(zip(columns, row)) for row in cur.fetchall()]
    finally:
        if conn is not None:
            conn.close()


def normalized_status(value):
    return str(value or "Pending").strip().casefold()


def normalized_type(value):
    return str(value or "").strip().casefold()


def compute_summary(rows):
    statuses = {status: 0 for status in EXPECTED_STATUSES}
    deployment_types = {"Automation": 0, "Manual": 0}
    successful_by_type = {"Automation": 0, "Manual": 0}
    failed_by_type = {"Automation": 0, "Manual": 0}

    for row in rows:
        status_key = normalized_status(row.get("status"))
        display_status = next(
            (status for status in EXPECTED_STATUSES if status.casefold() == status_key),
            str(row.get("status") or "Pending").strip() or "Pending",
        )
        statuses[display_status] = statuses.get(display_status, 0) + 1

        raw_type = normalized_type(row.get("deployment_type"))
        display_type = "Automation" if raw_type == "automation" else "Manual" if raw_type == "manual" else None
        if display_type:
            deployment_types[display_type] += 1
            if status_key == "successful":
                successful_by_type[display_type] += 1
            elif status_key == "failed":
                failed_by_type[display_type] += 1

    return {
        "total": len(rows),
        "successful": statuses.get("Successful", 0),
        "failed": statuses.get("Failed", 0),
        "automation": deployment_types["Automation"],
        "manual": deployment_types["Manual"],
        "successful_automation": successful_by_type["Automation"],
        "successful_manual": successful_by_type["Manual"],
        "failed_automation": failed_by_type["Automation"],
        "failed_manual": failed_by_type["Manual"],
        "statuses": statuses,
        "deployment_types": deployment_types,
    }


def excel_value(value):
    """Convert DB values to clean, Excel-friendly values."""
    if value is None:
        return ""
    if isinstance(value, datetime):

        if value.tzinfo is not None:
            return value.astimezone(REPORT_TZ).replace(tzinfo=None)
        return value
    if isinstance(value, (date, time)):
        return value
    return value


def style_header(ws, row_num, start_col, end_col):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="D9EAF7"))


def build_excel_report(report_date, rows, output_path):
    summary = compute_summary(rows)
    wb = Workbook()
    overview = wb.active
    overview.title = "Daily Summary"
    detail = wb.create_sheet("Deployment Details")

    overview.merge_cells("A1:D1")
    overview["A1"] = "DAILY DEPLOYMENT REPORT"
    overview["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    overview["A1"].fill = PatternFill("solid", fgColor="17365D")
    overview["A1"].alignment = Alignment(vertical="center")
    overview.row_dimensions[1].height = 30
    overview["A2"] = "Report Date"
    overview["B2"] = report_date
    overview["B2"].number_format = "dd-mmm-yyyy"
    overview["C2"] = "Generated At (IST)"
    overview["D2"] = datetime.now(REPORT_TZ).replace(tzinfo=None)
    overview["D2"].number_format = "dd-mmm-yyyy hh:mm:ss AM/PM"

    overview["A4"] = "Metric"
    overview["B4"] = "Count"
    style_header(overview, 4, 1, 2)
    metrics = [
        ("Total deployments scheduled for this date", summary["total"]),
        ("Successful deployments (completed successfully)", summary["successful"]),
        ("Failed deployments", summary["failed"]),
        ("Automation deployments", summary["automation"]),
        ("Manual deployments", summary["manual"]),
        ("Successful Automation deployments", summary["successful_automation"]),
        ("Successful Manual deployments", summary["successful_manual"]),
        ("Failed Automation deployments", summary["failed_automation"]),
        ("Failed Manual deployments", summary["failed_manual"]),
    ]
    for row_num, (label, count) in enumerate(metrics, start=5):
        overview.cell(row_num, 1, label)
        overview.cell(row_num, 2, count)
        overview.cell(row_num, 1).alignment = Alignment(wrap_text=True, vertical="top")

    status_header_row = 16
    overview.cell(status_header_row, 1, "Status Breakdown")
    overview.cell(status_header_row, 2, "Count")
    style_header(overview, status_header_row, 1, 2)
    for row_num, status in enumerate(EXPECTED_STATUSES, start=status_header_row + 1):
        overview.cell(row_num, 1, status)
        overview.cell(row_num, 2, summary["statuses"].get(status, 0))

    type_start = 16
    overview.cell(type_start, 4, "Deployment Type Breakdown")
    overview.cell(type_start, 5, "Count")
    style_header(overview, type_start, 4, 5)
    overview.cell(type_start + 1, 4, "Automation")
    overview.cell(type_start + 1, 5, summary["deployment_types"]["Automation"])
    overview.cell(type_start + 2, 4, "Manual")
    overview.cell(type_start + 2, 5, summary["deployment_types"]["Manual"])

    overview.freeze_panes = "A5"
    overview.column_dimensions["A"].width = 46
    overview.column_dimensions["B"].width = 15
    overview.column_dimensions["C"].width = 22
    overview.column_dimensions["D"].width = 25
    overview.column_dimensions["E"].width = 14
    for row in overview.iter_rows(min_row=4):
        for cell in row:
            if cell.value is not None and cell.row not in (4, status_header_row):
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    detail.append(DETAIL_HEADERS)
    style_header(detail, 1, 1, len(DETAIL_HEADERS))
    detail.row_dimensions[1].height = 34

    data_keys = [
        "id", "cr_number", "rfid", "deployment_date", "slot_start", "slot_end",
        "application_name", "deployment_type", "change_type", "status",
        "change_requestor", "released_by", "verified_and_deployed_by",
        "incident_number", "roll_back", "last_modified_by_email", "last_modified_at",
        "release_notes", "comments",
    ]
    for row in rows:
        detail.append([excel_value(row.get(key)) for key in data_keys])

    detail.freeze_panes = "A2"
    detail.auto_filter.ref = f"A1:S{max(1, detail.max_row)}"
    widths = {
        "A": 11, "B": 16, "C": 13, "D": 16, "E": 14, "F": 14,
        "G": 28, "H": 16, "I": 20, "J": 15, "K": 24, "L": 22,
        "M": 25, "N": 18, "O": 12, "P": 32, "Q": 23, "R": 38, "S": 38,
    }
    for col, width in widths.items():
        detail.column_dimensions[col].width = width

    for row_num in range(2, detail.max_row + 1):
        detail.cell(row_num, 4).number_format = "dd-mmm-yyyy"
        detail.cell(row_num, 5).number_format = "hh:mm AM/PM"
        detail.cell(row_num, 6).number_format = "hh:mm AM/PM"
        detail.cell(row_num, 17).number_format = "dd-mmm-yyyy hh:mm:ss AM/PM"
        for cell in detail[row_num]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if rows:
        table = Table(displayName="DailyDeployments", ref=f"A1:S{detail.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        detail.add_table(table)

    wb.save(output_path)
    return summary


def send_report_email(report_date, report_path, summary, recipients):
    """Send this report only to the explicitly configured manager list."""
    host = (os.getenv("MAIL_SERVER") or "").strip()
    username = (os.getenv("MAIL_USERNAME") or "").strip()
    password = os.getenv("MAIL_PASSWORD", "")
    from_email = (os.getenv("MAIL_FROM_EMAIL") or username).strip()
    from_name = (os.getenv("MAIL_FROM_NAME") or "Deployment Tracker").strip()

    missing = []
    if not host:
        missing.append("MAIL_SERVER")
    if not username:
        missing.append("MAIL_USERNAME")
    if not password:
        missing.append("MAIL_PASSWORD")
    if not from_email:
        missing.append("MAIL_FROM_EMAIL")
    if missing:
        raise RuntimeError("Missing SMTP setting(s) in .env: " + ", ".join(missing))

    port = int(os.getenv("MAIL_PORT", "587"))
    use_ssl = os.getenv("MAIL_USE_SSL", "false").strip().lower() in ("1", "true", "yes", "on")
    use_tls = os.getenv("MAIL_USE_TLS", "true").strip().lower() in ("1", "true", "yes", "on")
    if use_ssl and use_tls:
        raise RuntimeError("Set only one of MAIL_USE_SSL=true or MAIL_USE_TLS=true, not both.")
    timeout = int(os.getenv("MAIL_TIMEOUT_SECONDS", "30"))

    subject = f"Daily Deployment Report - {report_date.strftime('%d-%b-%Y')}"
    body = "\n".join([
        f"Daily Deployment Tracker report for {report_date.strftime('%d-%b-%Y')} (India time).",
        "",
        f"Total deployment records: {summary['total']}",
        f"Successful: {summary['successful']}",
        f"Failed: {summary['failed']}",
        f"Automation: {summary['automation']} (successful {summary['successful_automation']}, failed {summary['failed_automation']})",
        f"Manual: {summary['manual']} (successful {summary['successful_manual']}, failed {summary['failed_manual']})",
        "",
        "The attached Excel workbook contains the status breakdown and individual deployment records.",
        #"The 'Last Modified By Email' column uses the latest email recorded in the deployment audit history.",
        #"This report was sent only to addresses configured in MANAGER_EMAILS.",
    ])

    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = formataddr((from_name, from_email))
    message["To"] = ", ".join(recipients)
    message.set_content(body)
    with open(report_path, "rb") as workbook_file:
        message.add_attachment(
            workbook_file.read(),
            maintype="application",
            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=report_path.name,
        )

    if use_ssl:
        with smtplib.SMTP_SSL(
            host, port, timeout=timeout, context=ssl.create_default_context()
        ) as smtp:
            smtp.login(username, password)
            smtp.send_message(message, to_addrs=recipients)
    else:
        with smtplib.SMTP(host, port, timeout=timeout) as smtp:
            smtp.ehlo()
            if use_tls:
                smtp.starttls(context=ssl.create_default_context())
                smtp.ehlo()
            smtp.login(username, password)
            smtp.send_message(message, to_addrs=recipients)


def main():
    report_date = datetime.now(REPORT_TZ).date()
    recipients = get_manager_recipients()
    logger.info("Building daily report for %s; manager recipient count=%d", report_date, len(recipients))

    rows = fetch_daily_deployments(report_date)
    output_path = REPORT_DIR / f"Deployment_Report_{report_date.isoformat()}.xlsx"
    summary = build_excel_report(report_date, rows, output_path)
    logger.info(
        "Excel report built: %s | total=%d successful=%d failed=%d automation=%d manual=%d",
        output_path, summary["total"], summary["successful"], summary["failed"],
        summary["automation"], summary["manual"],
    )

    send_report_email(report_date, output_path, summary, recipients)
    logger.info("Daily report emailed successfully to the configured manager recipient(s) only.")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        logger.exception("Daily Deployment Report failed")
        raise
